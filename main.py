import os
import re
import logging
from pathlib import Path
from typing import Dict, Optional
import pandas as pd

import yaml
import json
import pytesseract
from pdf2image import convert_from_path
from dotenv import load_dotenv
from openai import OpenAI
from openpyxl import load_workbook
from PIL import Image

# Setup logging
logging.basicConfig(
    level=logging.INFO, format="%(asctime)s - %(levelname)s - %(message)s"
)
logging.getLogger("httpx").setLevel(logging.WARNING)
logging.getLogger("openai").setLevel(logging.WARNING)


# load .env file if present
try:
    load_dotenv(".env")
    logging.debug("Loaded environment variables from .env file.")
except ImportError:
    logging.warning("dotenv module not found, skipping environment variable loading.")

# Load settings from YAML
with open("settings.yaml", "r", encoding="utf-8") as f:
    settings = yaml.safe_load(f)

INPUT_PATH = Path(settings.get("input_path", "input"))
OUTPUT_PATH = Path(settings.get("output_path", "output"))
INPUT_EXCEL = INPUT_PATH / settings.get("input_excel", None)
OUTPUT_EXCEL = OUTPUT_PATH / settings.get("output_excel", None)
INVOICE_DIR = INPUT_PATH / settings.get("invoice_directory", None)
EXTRACTION_METHOD = settings.get("extraction_method")  # Options: "ocr_gpt"
TESSERACT_PATH = settings.get("tesseract_path")
OPENAI_API_KEY = os.getenv("OPENAI_API_KEY")


# Configure Tesseract executable path if required
if TESSERACT_PATH:
    pytesseract.pytesseract.tesseract_cmd = TESSERACT_PATH

# Configure OpenAI API key if required
if OPENAI_API_KEY:
    openai_client = OpenAI(api_key=OPENAI_API_KEY)


class InvoiceMatcher:

    def __init__(
        self, input_file: Path, invoice_dir: Path, output_dir: Path, method: str
    ):
        self.input_file = input_file
        self.invoice_dir = invoice_dir
        self.output_dir = output_dir
        self.output_file = OUTPUT_EXCEL
        self.method = method

        self.data = self.load_input()
        self.extracted_data = []
        self.processed_rows = []
        self.uncertain_matches = []
        self.unmatched_invoices = []

        # Check if output file exists and has the required sheets
        if not self.output_file.exists():
            logging.warning(
                f"Output file {self.output_file} does not exist. Creating a new one."
            )
            self.wb = None
        else:
            logging.info(f"Loading existing output file: {self.output_file}")
            self.wb = load_workbook(self.output_file)
            required_sheets = {"Matches", "Uncertain"}
            available_sheets = set(self.wb.sheetnames)

            self.output_excel_match = required_sheets.issubset(available_sheets)
            if not self.output_excel_match:
                logging.warning(
                    f"Output file {self.output_file} does not contain all required sheets: {required_sheets}"
                )

        # Make all directories if they do not exist
        self.output_dir.mkdir(parents=True, exist_ok=True)

    def update_settings(self):
        logging.debug("Updating settings.yaml")
        with open("settings.yaml", "w", encoding="utf-8") as f:
            yaml.dump(settings, f)

    def load_input(self) -> pd.DataFrame:
        logging.info(f"Loading CSV file from {self.input_file}")
        logging.warning("This script is not idempotent. Re-running it on the same input file will result in duplicate entries and incorrect invoice numbers.")
        try:
            df = pd.read_csv(self.input_file, sep=";", encoding="utf-8", decimal=",")
            df.columns = df.columns.str.strip().str.lower()
            logging.debug(f"Loaded columns: {df.columns.tolist()}")
            return df
        except FileNotFoundError:
            logging.warning(
                f"Input file {self.input_file} not found. All invoices will be processed as unmatched."
            )
            return pd.DataFrame()
        except Exception as e:
            logging.error(f"Failed to load input file: {e}")
            return pd.DataFrame()

    def extract_invoice_data(self, invoice_file_path: Path) -> Optional[Dict[str, str]]:
        if self.method == "ocr_gpt":
            return self.extract_ocr_gpt(invoice_file_path)
        else:
            logging.error(f"Unknown extraction method: {self.method}")
            return None

    def extract_ocr_gpt(self, invoice_file_path: Path) -> Optional[Dict[str, str]]:
        logging.info(f"Extracting (OCR + GPT) from: {invoice_file_path}")
        logging.debug(
            f"Using Tesseract from: {pytesseract.pytesseract.tesseract_cmd}, version: {pytesseract.get_tesseract_version()}"
        )

        if (
            self.output_file.exists()  # Check if output file exists
            and not settings.get(
                "process_all", False
            )  # If not overriding existing results
            and self.output_excel_match  # Checked on init if file has all sheets needed
        ):
            logging.debug(
                f"Checking output file: {self.output_file} for existing results for {invoice_file_path}"
            )
            df_output_matches = pd.read_excel(
                self.output_file, sheet_name="Matches", engine="openpyxl"
            )
            df_output_uncertain = pd.read_excel(
                self.output_file, sheet_name="Uncertain", engine="openpyxl"
            )
            df_output = pd.concat(
                [df_output_matches, df_output_uncertain], ignore_index=True
            )

            this_output = df_output[df_output["file_path"] == str(invoice_file_path)]
            if not this_output.empty:
                logging.info(
                    f"All information already extracted skipping extraction. Use process_all=True in settings.yaml to force re-extraction."
                )
                return {
                    "invoice_number": this_output["invoice_number"].iloc[0],
                    "total_amount": this_output["totaal bedrag"].iloc[0],
                    "receiver": this_output["iban betaald"].iloc[0],
                    "file_path": str(invoice_file_path),
                }

        try:
            if invoice_file_path.suffix.lower() == ".pdf":
                logging.debug(f"Converting PDF to images for OCR: {invoice_file_path}")
                images = convert_from_path(str(invoice_file_path))
            elif invoice_file_path.suffix.lower() in [".jpg", ".jpeg", ".png"]:
                logging.debug(f"Using image directly for OCR: {invoice_file_path}")
                image = Image.open(invoice_file_path)
                images = [image]
            else:
                logging.error(f"Unsupported file type: {invoice_file_path.suffix}")
                return {
                    "invoice_number": "onbekend",
                    "total_amount": "onbekend",
                    "receiver": "onbekend",
                    "file_path": str(invoice_file_path),
                }

            text = "".join(pytesseract.image_to_string(image) for image in images)
            text = re.sub(r" ", "", text).lower()

            prompt = """Lees de inhoud van de factuur of bon en extraheer de volgende gegevens:
                        - Factuurnummer (bijv. factuurnummer, invoice number)
                        - Totaal bedrag (vaak onderaan, bijv. totaal, totaal bedrag, total amount, factuurbedrag)
                        - Ontvanger (IBAN), die volgt patroon [a-z]{2}\d{2}[a-z0-9]{4}\d{10}

                        Geef de resultaten terug in JSON-formaat met de volgende structuur:
                        1. invoice_number (Factuurnummer)
                        2. total_amount (Totaal bedrag)
                        3. receiver (IBAN)
                        Bij missende waarden, gebruik null

                        Getallen kunnen komma's bevatten, dus converteer ze naar een float.
                    """

            response = openai_client.chat.completions.create(
                model="gpt-4.1-mini",
                messages=[
                    {
                        "role": "system",
                        "content": """
                            Jij bent een expert in het extraheren van factuurgegevens uit PDF's.
                            Je bent erg secuur en neemt de tijd om een goede en gevalideerder JSON als output te geven.
                            Je reageert in plain text, zonder markdown of andere opmaak. Dus geen ```json of dergelijke.
                            Neem de tijd om de PDF goed te lezen en de juiste gegevens te extraheren.
                            Kijk goed of je een IBAN kunt vinden, en of het totaal bedrag een getal is.
                            """,
                    },
                    {
                        "role": "user",
                        "content": f"""{prompt} \n\n(PDF contents extracted via OCR) \n{text}""",
                    },
                ],
            )
        except Exception as e:
            logging.error(f"Error during OCR + GPT extraction: {e}")
            return None

        usage = response.usage
        logging.debug(
            f"Prompt tokens: {usage.prompt_tokens}, Completion tokens: {usage.completion_tokens}, Total: {usage.total_tokens}"
        )

        result = response.choices[0].message.content
        data = json.loads(result)
        logging.info(f"Extracted data: {data}")

        return {
            "invoice_number": data.get("invoice_number"),
            "total_amount": (
                float(data.get("total_amount")) if data.get("total_amount") else None
            ),
            "receiver": data.get("receiver"),
            "file_path": str(invoice_file_path),
        }

    def extract_all_data(self):
        logging.info("Extracting data from all PDFs...")
        for pdf_file in self.invoice_dir.iterdir():
            invoice_data = self.extract_invoice_data(pdf_file)
            if invoice_data:
                self.extracted_data.append(invoice_data)

    def compare_and_update(self):
        logging.info("Comparing input data with extracted invoices...")
        iban_col = "iban betaald"
        amount_col = "totaal bedrag"

        if iban_col not in self.data.columns or amount_col not in self.data.columns:
            logging.error(f"Required columns '{iban_col}' or '{amount_col}' not found in input file.")
            return

        available_invoices = self.extracted_data.copy()

        invoice_number = settings.get("last_invoice_number", 0)

        for index, row in self.data.iterrows():
            invoice_number += 1
            invoice_number_styled = f"F{invoice_number:04d}"

            output_row = row.to_frame().T.copy()
            output_row.insert(0, "ID", invoice_number_styled)

            row_iban = row.get(iban_col)
            row_amount = row.get(amount_col)

            matching_invoices = []
            if pd.notna(row_iban) and pd.notna(row_amount):
                for invoice in available_invoices:
                    invoice_iban = invoice.get("receiver")
                    invoice_amount = invoice.get("total_amount")

                    iban_match = (
                        isinstance(row_iban, str) and
                        isinstance(invoice_iban, str) and
                        row_iban.lower() == invoice_iban.lower()
                    )

                    try:
                        amount_match = pd.notna(invoice_amount) and abs(float(row_amount) - float(invoice_amount)) < 0.01
                    except (ValueError, TypeError):
                        amount_match = False

                    if iban_match and amount_match:
                        matching_invoices.append(invoice)

            if len(matching_invoices) == 1:
                matched_invoice = matching_invoices[0]
                logging.info(f"Single match found for row {index} with invoice {matched_invoice['file_path']}")

                output_row["file_path"] = matched_invoice["file_path"]
                output_row["invoice_number"] = matched_invoice.get("invoice_number")

                try:
                    new_file_name = f"{invoice_number_styled}{Path(matched_invoice['file_path']).suffix}"
                    new_file_path = self.invoice_dir / new_file_name
                    Path(matched_invoice["file_path"]).rename(new_file_path)
                    logging.info(f"Renamed invoice file to {new_file_name}")
                except Exception as e:
                    logging.error(f"Failed to rename invoice file: {e}")
                    output_row["file_path"] = "<rename_failed>"

                available_invoices.remove(matched_invoice)
                self.processed_rows.append(output_row)

            elif len(matching_invoices) > 1:
                logging.warning(f"Multiple matches for row {index}")
                output_row["file_path"] = ", ".join([inv["file_path"] for inv in matching_invoices])
                output_row["invoice_number"] = ", ".join([str(inv.get("invoice_number", '')) for inv in matching_invoices])
                self.uncertain_matches.append(output_row)
                for inv in matching_invoices:
                    if inv in available_invoices:
                        available_invoices.remove(inv)
            else: # No match
                logging.warning(f"No match found for row {index}")
                output_row["file_path"] = ""
                output_row["invoice_number"] = ""
                self.processed_rows.append(output_row)

        settings["last_invoice_number"] = invoice_number

        self.unmatched_invoices = available_invoices

    def save_results(self):
        logging.info("Saving results to Excel...")
        with pd.ExcelWriter(OUTPUT_EXCEL, engine="xlsxwriter") as writer:
            if self.processed_rows:
                pd.concat(self.processed_rows).to_excel(
                    writer, sheet_name="Matches", index=False
                )
            if self.uncertain_matches:
                pd.concat(self.uncertain_matches).to_excel(
                    writer, sheet_name="Uncertain", index=False
                )
            if self.unmatched_invoices:
                pd.DataFrame(self.unmatched_invoices).to_excel(
                    writer, sheet_name="No Match", index=False
                )


if __name__ == "__main__":
    matcher = InvoiceMatcher(INPUT_EXCEL, INVOICE_DIR, OUTPUT_PATH, EXTRACTION_METHOD)
    matcher.extract_all_data()
    matcher.compare_and_update()
    matcher.save_results()
    matcher.update_settings()
    logging.info("Invoice matching completed.")
